"""Self-contained deep audit suite for facebook-extractor.

Everything runs against a SCRATCH COPY of the database (or a fresh synthetic
one), so production data is never modified.  ``DEEPSEEK_API_KEY`` is forced
empty: the suite is fully offline and never calls the AI endpoint.

Usage (from the project root):
    python tests/audit_suite.py

Exit code 0 = all checks passed.  A scratch directory is printed at the end.
"""

from __future__ import annotations

import csv
import json
import os
import shutil
import sys
import tempfile
from pathlib import Path
from unittest import mock

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))

# --- scratch environment (must happen before importing config) --------------
scratch_dir = Path(tempfile.mkdtemp(prefix="fbx_audit_"))
scratch_db = scratch_dir / "audit.db"
real_db = PROJECT / "data" / "facebook_content.db"
# The suite must pass with or without a real database present: CI has none.
# Set AUDIT_SYNTHETIC_ONLY=1 to force the no-source-database mode locally.
if real_db.exists() and os.environ.get("AUDIT_SYNTHETIC_ONLY") != "1":
    shutil.copy2(real_db, scratch_db)
os.environ["DATABASE_PATH"] = str(scratch_db)
os.environ["DEEPSEEK_API_KEY"] = ""  # keep the suite fully offline

from ai.deepseek_client import DeepSeekError, parse_json_response  # noqa: E402
from config import load_settings  # noqa: E402
from database import (  # noqa: E402
    counts_by,
    db_session,
    existing_statuses,
    failed_urls,
    get_item,
    get_item_by_url,
    get_transcript_segments,
    init_db,
    list_items,
    list_items_page,
    replace_transcript_segments,
    save_summary,
    search_items,
    stale_summary_item_ids,
    summaries_for_ids,
    update_item_fields,
    upsert_item,
)
from export.csv_exporter import export_csv  # noqa: E402
from export.json_exporter import export_json  # noqa: E402
from export.markdown_exporter import export_markdown  # noqa: E402
from extractor import net, pipeline  # noqa: E402
from extractor.classifier import classify  # noqa: E402
from models import (  # noqa: E402
    ExtractedItem,
    ExtractionOutcome,
    SummaryPayload,
    TranscriptSegment,
)

settings = load_settings()
assert str(settings.database_path) == str(scratch_db), (
    f"scratch DB not active: {settings.database_path}"
)

results: list[tuple[str, bool, str]] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    ok = bool(condition)
    results.append((name, ok, detail))
    suffix = f"  [{detail}]" if detail and not ok else ""
    print(f"{'PASS' if ok else 'FAIL'}  {name}{suffix}")


def fail_outcome(message: str, login: bool = False) -> ExtractionOutcome:
    outcome = ExtractionOutcome()
    outcome.error = message
    outcome.needs_login = login
    return outcome


# ---------------------------------------------------------------------------
# Deterministic fixtures (created inside the scratch DB only)
# ---------------------------------------------------------------------------
init_db(settings.database_path)

REEL_URL = "https://www.facebook.com/share/r/AUDITFIXTURE1/"
POST_URL = "https://www.facebook.com/share/p/AUDITFIXTURE2/"
MANUAL_URL = "https://www.facebook.com/share/p/AUDITFIXTURE3/"

whisper_item = ExtractedItem(
    source_url=REEL_URL,
    content_type="reel",
    status="extracted",
    title="fixture reel",
    caption="fixture caption for the reel " * 10,
    transcript_text="fixture whisper transcript words " * 10,
    transcript_source="whisper",
    transcript_language="en",
    view_count=1000,
)
whisper_id = upsert_item(whisper_item, settings.database_path)
replace_transcript_segments(
    whisper_id,
    [
        TranscriptSegment(start_seconds=0.0, end_seconds=2.5, text="fixture segment one"),
        TranscriptSegment(start_seconds=2.5, end_seconds=5.0, text="fixture segment two"),
    ],
    settings.database_path,
)

partial_item = ExtractedItem(
    source_url=POST_URL,
    content_type="post",
    status="partial",
    caption="truncated preview text...",
    transcript_text="truncated preview text...",
    transcript_source="description",
    needs_manual_import=True,
)
upsert_item(partial_item, settings.database_path)

manual_item = ExtractedItem(
    source_url=MANUAL_URL,
    content_type="post",
    status="extracted",
    caption="manually imported caption " * 10,
    transcript_text="manually imported transcript " * 10,
    transcript_source="manual",
)
manual_id = upsert_item(manual_item, settings.database_path)
replace_transcript_segments(
    manual_id,
    [TranscriptSegment(start_seconds=0.0, end_seconds=None, text="manually imported transcript")],
    settings.database_path,
)

print("=== 1. URL classifier ===")
expected_types = {
    "https://www.facebook.com/share/p/ABC123/": "post",
    "https://www.facebook.com/share/r/ABC123/": "reel",
    "https://www.facebook.com/share/v/ABC123/": "video",
    "https://www.facebook.com/": "homepage",
    "https://www.facebook.com/share/p/AAA/share/r/BBB/": "malformed",
    "https://example.com/x": "unknown",
    "https://www.facebook.com/reel/123456": "reel",
    "https://www.facebook.com/watch/?v=123": "video",
    "": "malformed",
}
for url, expected in expected_types.items():
    got = classify(url).content_type.value
    check(f"classify {expected!r}", got == expected, f"got {got!r}")

repaired = classify("https://www.facebook.com/share/p/AAA/share/r/BBB/")
check(
    "malformed repair candidate first",
    repaired.candidates[0] == "https://www.facebook.com/share/p/AAA/",
    str(repaired.candidates),
)
check(
    "malformed candidates include the original",
    "https://www.facebook.com/share/p/AAA/share/r/BBB/" in repaired.candidates,
)

print("\n=== 2. HTTP retry logic (429 / 5xx / Retry-After) ===")


class FakeResponse:
    def __init__(self, status, headers=None):
        self.status_code = status
        self.headers = headers or {}
        self.url = "http://fake"


class FlakySession:
    def __init__(self, statuses):
        self.statuses = list(statuses)
        self.calls = 0

    def request(self, *args, **kwargs):
        self.calls += 1
        return FakeResponse(self.statuses.pop(0) if self.statuses else 200)


flaky = FlakySession([429, 500, 200])
with mock.patch.object(net.time, "sleep", lambda seconds: None):
    response = net.request_with_retries(flaky, "GET", "http://fake", settings)
check(
    "429+500 retried then success",
    response.status_code == 200 and flaky.calls == 3,
    f"calls={flaky.calls} status={response.status_code}",
)

always_429 = FlakySession([429] * 6)
with mock.patch.object(net.time, "sleep", lambda seconds: None):
    response = net.request_with_retries(always_429, "GET", "http://fake", settings)
check(
    "exhausted 429 returns the last response",
    response.status_code == 429 and always_429.calls == max(1, settings.max_retries),
    f"calls={always_429.calls} status={response.status_code}",
)

delays: list[float] = []
with mock.patch.object(net.time, "sleep", lambda seconds: delays.append(seconds)):
    net.request_with_retries(FlakySession([429, 200]), "GET", "http://fake", settings)
check("backoff sleep happens between retries", len(delays) == 1, str(delays))

retry_after_delay = net._retry_delay(
    settings, attempt=0, response=FakeResponse(429, {"Retry-After": "7"})
)
check("Retry-After honoured (>=7s)", retry_after_delay >= 7.0, f"{retry_after_delay}")

def _no_redirect(url, *args, **kwargs):
    """Keep the suite hermetic: never touch the network for redirect checks."""
    return url, 200, None


print("\n=== 3. Data preservation on failed refresh ===")
before = get_item_by_url(MANUAL_URL, settings.database_path)
with mock.patch.object(
    pipeline.ytdlp_extractor, "extract_video", lambda *a, **k: fail_outcome("simulated")
), mock.patch.object(
    pipeline.facebook_post_extractor, "extract_post", lambda *a, **k: fail_outcome("simulated")
), mock.patch.object(pipeline, "resolve_redirects", _no_redirect):
    record = pipeline.process_url(MANUAL_URL, settings, transcribe=False, summarize=False)
after = get_item_by_url(MANUAL_URL, settings.database_path)
check("preserved_previous flag set", record.get("preserved_previous") is True)
check(
    "caption preserved",
    len(after["caption"] or "") == len(before["caption"] or ""),
    f"{len(after['caption'] or '')} vs {len(before['caption'] or '')}",
)
check(
    "transcript preserved with its source",
    after["transcript_source"] == "manual"
    and len(after["transcript_text"] or "") == len(before["transcript_text"] or ""),
)
check("status not downgraded", after["status"] == before["status"], after["status"])
check(
    "fresh failure noted in error_message",
    "kept previous data" in (after["error_message"] or ""),
)

print("\n=== 4. Transcript rank protection (whisper not downgraded) ===")
w_before = get_item_by_url(REEL_URL, settings.database_path)
w_segments_before = len(get_transcript_segments(w_before["id"], settings.database_path))

weaker = ExtractionOutcome()
weaker.fields = {"caption": "short fresh caption without any transcript", "resolved_url": REEL_URL}
with mock.patch.object(
    pipeline.ytdlp_extractor, "extract_video", lambda *a, **k: weaker
), mock.patch.object(
    pipeline.facebook_post_extractor, "extract_post", lambda *a, **k: fail_outcome("simulated")
), mock.patch.object(pipeline, "resolve_redirects", _no_redirect):
    record = pipeline.process_url(REEL_URL, settings, transcribe=False, summarize=False)
w_after = get_item_by_url(REEL_URL, settings.database_path)
w_segments_after = len(get_transcript_segments(w_after["id"], settings.database_path))
check("kept_previous flag set", record.get("kept_previous") is True)
check(
    "whisper transcript kept",
    w_after["transcript_source"] == "whisper"
    and len(w_after["transcript_text"] or "") == len(w_before["transcript_text"] or ""),
)
check(
    "segments kept",
    w_segments_after == w_segments_before and w_segments_after > 0,
    f"{w_segments_after} vs {w_segments_before}",
)

print("\n=== 5. Manual import semantics + stickiness ===")
p_before = get_item_by_url(POST_URL, settings.database_path)
check("fixture starts partial", p_before["status"] == "partial", p_before["status"])
long_text = "Manually pasted full post text. " * 20
pipeline.manual_import(POST_URL, settings, text=long_text)
p_after = get_item_by_url(POST_URL, settings.database_path)
check("caption replaced by longer manual text", (p_after["caption"] or "").startswith("Manually pasted"))
check("transcript source manual", p_after["transcript_source"] == "manual")
check(
    "status extracted + manual flag cleared",
    p_after["status"] == "extracted" and not p_after["needs_manual_import"],
)

truncated = ExtractionOutcome()
truncated.fields = {"caption": "truncated repost preview...", "resolved_url": POST_URL}
truncated.warnings.append(
    "Facebook only served a truncated public preview for this post; the full text may require manual import"
)
with mock.patch.object(
    pipeline.facebook_post_extractor, "extract_post", lambda *a, **k: truncated
), mock.patch.object(
    pipeline.ytdlp_extractor, "extract_video", lambda *a, **k: fail_outcome("n/a")
), mock.patch.object(pipeline, "resolve_redirects", _no_redirect):
    record = pipeline.process_url(POST_URL, settings, transcribe=False, summarize=False)
p_final = get_item_by_url(POST_URL, settings.database_path)
check(
    "manual data sticky against truncated refresh",
    record.get("kept_previous") is True
    and (p_final["caption"] or "").startswith("Manually pasted")
    and p_final["status"] == "extracted",
)

print("\n=== 6. Exports (CSV / Markdown / JSON) ===")
poison = ExtractedItem(
    source_url="https://www.facebook.com/audit/poison",
    content_type="post",
    status="extracted",
    title="@SUM(1+1)",
    caption='=HYPERLINK("http://evil.example","click")',
)
upsert_item(poison, settings.database_path)
total_items = len(list_items(settings.database_path))

csv_path = export_csv(settings, scratch_dir / "export_items.csv")
with csv_path.open(encoding="utf-8-sig", newline="") as handle:
    rows = list(csv.DictReader(handle))
check("CSV row count matches DB", len(rows) == total_items, f"{len(rows)} vs {total_items}")
check("CSV has summary columns", "english_summary" in rows[0] and "source_url" in rows[0])
poison_row = next(r for r in rows if r["source_url"].endswith("/audit/poison"))
check("CSV formula guard on caption", poison_row["caption"].startswith("'="), poison_row["caption"][:20])
check("CSV formula guard on title", poison_row["title"].startswith("'@"), poison_row["title"][:20])
check(
    "CSV numeric fields untouched",
    all(row["duration_seconds"] == "" or not row["duration_seconds"].startswith("'") for row in rows),
)

md_text = export_markdown(settings, scratch_dir / "report.md").read_text(encoding="utf-8")
check("Markdown grouped by content type", "## Videos" in md_text or "## Reels" in md_text)
check("Markdown contains AI Summary sections", "AI Summary" in md_text)

payload = json.loads(export_json(settings, scratch_dir / "export_items.json").read_text(encoding="utf-8"))
check("JSON export item count", payload["item_count"] == total_items, str(payload["item_count"]))
check(
    "JSON export includes transcript segments",
    any(entry.get("transcript_segments") for entry in payload["items"]),
)

print("\n=== 7. Database integrity invariants ===")
with db_session(settings.database_path) as conn:
    orphan_summaries = conn.execute(
        "SELECT COUNT(*) AS c FROM summaries s LEFT JOIN items i ON i.id = s.item_id WHERE i.id IS NULL"
    ).fetchone()["c"]
    orphan_segments = conn.execute(
        "SELECT COUNT(*) AS c FROM transcript_segments t LEFT JOIN items i ON i.id = t.item_id WHERE i.id IS NULL"
    ).fetchone()["c"]
    empty_transcripts = conn.execute(
        "SELECT COUNT(*) AS c FROM items WHERE transcript_source IN ('subtitles','whisper','manual') "
        "AND (transcript_text IS NULL OR transcript_text = '')"
    ).fetchone()["c"]
    extracted_manual = conn.execute(
        "SELECT COUNT(*) AS c FROM items WHERE status = 'extracted' AND needs_manual_import = 1"
    ).fetchone()["c"]
    skipped_with_caption = conn.execute(
        "SELECT COUNT(*) AS c FROM items WHERE status = 'skipped' AND caption IS NOT NULL AND caption != ''"
    ).fetchone()["c"]
    manual_without_segments = conn.execute(
        "SELECT COUNT(*) AS c FROM items i WHERE i.transcript_source = 'manual' "
        "AND NOT EXISTS (SELECT 1 FROM transcript_segments t WHERE t.item_id = i.id)"
    ).fetchone()["c"]
    segments_for_no_transcript = conn.execute(
        "SELECT COUNT(*) AS c FROM transcript_segments t JOIN items i ON i.id = t.item_id "
        "WHERE i.transcript_source = 'none'"
    ).fetchone()["c"]
check("no orphan summaries", orphan_summaries == 0, str(orphan_summaries))
check("no orphan transcript segments", orphan_segments == 0, str(orphan_segments))
check("no real transcript source with empty text", empty_transcripts == 0, str(empty_transcripts))
check("extracted items never need manual import", extracted_manual == 0, str(extracted_manual))
check("skipped items carry no caption", skipped_with_caption == 0, str(skipped_with_caption))
check("manual transcript always has segments", manual_without_segments == 0, str(manual_without_segments))
check("no segments for transcript-less items", segments_for_no_transcript == 0, str(segments_for_no_transcript))

print("\n=== 8. SQL-injection guards ===")
try:
    update_item_fields(
        db_path=settings.database_path, item_id=1, **{"error_message; DROP TABLE items;--": "x"}
    )
    check("update_item_fields rejects unknown column", False, "no exception raised")
except ValueError:
    check("update_item_fields rejects unknown column", True)

try:
    counts_by("status; DROP TABLE items", settings.database_path)
    check("counts_by rejects non-whitelisted column", False, "no exception raised")
except ValueError:
    check("counts_by rejects non-whitelisted column", True)

print("\n=== 9. DeepSeek client offline units ===")
check("fenced JSON parsed", parse_json_response('```json\n{"a": 1}\n```') == {"a": 1})
check("prose-wrapped JSON parsed", parse_json_response('Sure! {"a": {"b": 2}} done') == {"a": {"b": 2}})
try:
    parse_json_response("not json at all")
    check("invalid JSON raises DeepSeekError", False, "no exception raised")
except DeepSeekError:
    check("invalid JSON raises DeepSeekError", True)

weird = SummaryPayload.from_model_output(
    {
        "short_title": " T ",
        "key_points": "a;b;c",
        "detected_topics": ["x", "y"],
        "sentiment": ["positive"],
        "hashtags": None,
        "call_to_action": ["Do it", "now"],
    }
)
check("SummaryPayload coerces 'a;b;c'", weird.key_points == ["a", "b", "c"], str(weird.key_points))
check("SummaryPayload coerces list sentiment", weird.sentiment == "positive", weird.sentiment)
check("SummaryPayload joins list CTA", weird.call_to_action == "Do it; now", str(weird.call_to_action))
check("SummaryPayload handles None hashtags", weird.hashtags == [])

print("\n=== 10. JSONL append ===")
probe = ExtractedItem(source_url="https://www.facebook.com/audit/jsonl", status="extracted", title="probe")
pipeline.append_jsonl(probe, settings)
lines = settings.jsonl_path.read_text(encoding="utf-8").splitlines()
last = json.loads(lines[-1])
check("JSONL record appended and parseable", last["source_url"].endswith("/audit/jsonl"))

print("\n=== 11. Search index, content hashes, backlog helpers ===")
search_reel = ExtractedItem(
    source_url="https://www.facebook.com/audit/search-reel",
    content_type="reel",
    status="extracted",
    title="Agent patterns that ship",
    caption="Three agent patterns for production.",
    transcript_text="Planner executor and reflection loops explained.",
    transcript_source="whisper",
    view_count=5000,
)
search_id = upsert_item(search_reel, settings.database_path)
search_post = ExtractedItem(
    source_url="https://www.facebook.com/audit/search-post",
    content_type="post",
    status="partial",
    title="Resibo tips",
    caption="Filing receipts for freelancers (resibo).",
    transcript_text="I-save ang resibo para sa taxes.",
    needs_manual_import=True,
)
post_id = upsert_item(search_post, settings.database_path)

hits = search_items("agent", settings.database_path)
check("search matches title/caption", any(int(h["id"]) == search_id for h in hits), str([h["id"] for h in hits]))
check(
    "search returns a snippet",
    any(h.get("transcript_snippet") or h.get("caption_snippet") for h in hits),
)
check(
    "search matches transcript (Filipino term)",
    [int(h["id"]) for h in search_items("resibo", settings.database_path)] == [post_id],
)
check(
    "search uses AND semantics",
    len(search_items("agent production", settings.database_path)) >= 1
    and search_items("agent resibo", settings.database_path) == [],
)
check(
    "search status filter",
    all(h["status"] == "partial" for h in search_items("resibo", settings.database_path, status="partial"))
    and search_items("resibo", settings.database_path, status="extracted") == [],
)
check("search tolerates FTS operators", search_items('" OR * NEAR( ', settings.database_path) == [])
check("search blank query returns nothing", search_items("   ", settings.database_path) == [])
check(
    "search prefix wildcard",
    any(int(h["id"]) == search_id for h in search_items("agent*", settings.database_path)),
)
check(
    "trigger keeps index in sync",
    any(int(h["id"]) == search_id for h in search_items("planner executor", settings.database_path)),
)

with db_session(settings.database_path) as conn:
    # Simulate a database created before search existed: no index, no triggers.
    # (Do not use FTS5 'delete-all' here: emptying an external-content index
    # while the triggers stay active corrupts it on the next incremental
    # delete — a real hazard worth documenting in the test itself.)
    conn.execute("DROP TRIGGER IF EXISTS items_fts_ai")
    conn.execute("DROP TRIGGER IF EXISTS items_fts_ad")
    conn.execute("DROP TRIGGER IF EXISTS items_fts_au")
    conn.execute("DROP TABLE IF EXISTS items_fts")
check("search unavailable without an index", search_items("agent", settings.database_path) == [])
init_db(settings.database_path)
check("init_db recreates and rebuilds a missing index", len(search_items("agent", settings.database_path)) >= 1)
upsert_item(search_reel, settings.database_path)
check(
    "triggers keep the rebuilt index live",
    len(search_items("reflection", settings.database_path)) >= 1,
)

row = get_item(search_id, settings.database_path)
check("content hash recorded on upsert", bool(row and row.get("content_hash")))
demo_summary = SummaryPayload(
    short_title="Agent patterns", english_summary="EN summary", filipino_summary="FIL buod"
)
save_summary(search_id, "test-model", demo_summary, settings.database_path)
row2 = get_item(search_id, settings.database_path)
check(
    "summarized hash recorded with the summary",
    row2["summarized_hash"] == row["content_hash"],
)
check("fresh summary is not stale", search_id not in stale_summary_item_ids(settings.database_path))
changed = search_reel.model_copy(update={"transcript_text": "Completely different transcript about taxes."})
upsert_item(changed, settings.database_path)
check("changed content marks the summary stale", search_id in stale_summary_item_ids(settings.database_path))
save_summary(search_id, "test-model", demo_summary, settings.database_path)
check("re-summarising clears staleness", search_id not in stale_summary_item_ids(settings.database_path))

update_item_fields(db_path=settings.database_path, item_id=post_id, status="needs_login")
check("failed_urls picks retryable statuses", search_post.source_url in failed_urls(settings.database_path))

# Dedicated fixture: the suite must not depend on pre-existing database rows.
retry_fixture = ExtractedItem(
    source_url="https://www.facebook.com/audit/retry-partial",
    content_type="post",
    status="partial",
    caption="truncated preview awaiting manual import...",
    needs_manual_import=True,
)
upsert_item(retry_fixture, settings.database_path)
status_map = existing_statuses(settings.database_path)
custom_partial = failed_urls(settings.database_path, ["partial"])
check(
    "failed_urls honours custom statuses",
    retry_fixture.source_url in custom_partial
    and search_post.source_url not in custom_partial
    and all(status_map.get(url) == "partial" for url in custom_partial),
    f"{len(custom_partial)} partial urls",
)
check(
    "existing_statuses maps urls to status",
    existing_statuses(settings.database_path).get(search_reel.source_url) == "extracted",
)

page1, total1 = list_items_page(settings.database_path, limit=2, offset=0, sort="id", descending=True)
page2, total2 = list_items_page(settings.database_path, limit=2, offset=2, sort="id", descending=True)
check(
    "list_items_page totals match",
    total1 == total2 == len(list_items(settings.database_path)),
    f"{total1} vs {total2}",
)
check(
    "list_items_page pages are disjoint",
    len(page1) == len(page2) == 2 and {i["id"] for i in page1}.isdisjoint({i["id"] for i in page2}),
)
sorted_rows, _ = list_items_page(settings.database_path, sort="view_count", descending=True, limit=3)
views = [i["view_count"] or 0 for i in sorted_rows]
check("sort by view_count descending", views == sorted(views, reverse=True), str(views))
check("summaries_for_ids", set(summaries_for_ids([search_id], settings.database_path)) == {search_id})

print("\n=== 12. XLSX export ===")
from openpyxl import load_workbook  # noqa: E402

from export.xlsx_exporter import export_xlsx  # noqa: E402

xlsx_path = export_xlsx(settings, scratch_dir / "export_items.xlsx")
workbook = load_workbook(xlsx_path)
check("xlsx sheets", workbook.sheetnames == ["Items", "Summaries"], str(workbook.sheetnames))
items_sheet = workbook["Items"]
check(
    "xlsx item row count",
    items_sheet.max_row == len(list_items(settings.database_path)) + 1,
    f"{items_sheet.max_row}",
)
headers = [cell.value for cell in items_sheet[1]]
check("xlsx headers", "source_url" in headers and "transcript_text" in headers)
source_col = headers.index("source_url") + 1
caption_col = headers.index("caption") + 1
poison_row = next(
    row_index
    for row_index in range(2, items_sheet.max_row + 1)
    if str(items_sheet.cell(row=row_index, column=source_col).value or "").endswith("/audit/poison")
)
check(
    "xlsx formula guard on caption",
    str(items_sheet.cell(row=poison_row, column=caption_col).value).startswith("'="),
    str(items_sheet.cell(row=poison_row, column=caption_col).value)[:20],
)
check("xlsx summaries sheet populated", workbook["Summaries"].max_row >= 2)

# ---------------------------------------------------------------------------
failed = [name for name, ok, _ in results if not ok]
print(f"\n{'=' * 60}")
print(f"AUDIT SUITE: {len(results) - len(failed)}/{len(results)} passed")
if failed:
    print("FAILED:")
    for name in failed:
        print("  -", name)
print(f"scratch dir: {scratch_dir}")
sys.exit(1 if failed else 0)

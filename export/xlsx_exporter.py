"""XLSX export: a two-sheet workbook (items + summaries).

Sheet "Items" carries one row per item (transcript truncated to Excel's cell
limit); sheet "Summaries" carries the structured AI output.  All string cells
pass the shared spreadsheet formula-injection guard.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from config import Settings
from database import items_with_summaries
from export.csv_exporter import safe_cell

logger = logging.getLogger(__name__)

# Excel's hard limit per cell is 32,767 characters.
_EXCEL_CELL_LIMIT = 32000

ITEM_SHEET_COLUMNS = [
    "id",
    "source_url",
    "resolved_url",
    "content_type",
    "content_id",
    "status",
    "title",
    "author",
    "page_name",
    "upload_date",
    "duration_seconds",
    "view_count",
    "like_count",
    "comment_count",
    "share_count",
    "language",
    "transcript_source",
    "transcript_language",
    "caption",
    "transcript_text",
    "hashtags",
    "mentions",
    "external_links",
    "needs_login",
    "needs_manual_import",
    "malformed",
    "error_message",
    "extracted_at",
    "summarized_at",
]

SUMMARY_SHEET_COLUMNS = [
    "item_id",
    "model",
    "short_title",
    "english_summary",
    "filipino_summary",
    "key_points",
    "topics",
    "hashtags",
    "entities",
    "sentiment",
    "content_warnings",
    "suggested_tags",
    "virality_signals",
    "call_to_action",
    "created_at",
]


def _cell_value(value: Any) -> Any:
    if isinstance(value, list):
        value = "; ".join(str(item) for item in value)
    if isinstance(value, str):
        if len(value) > _EXCEL_CELL_LIMIT:
            value = value[:_EXCEL_CELL_LIMIT] + " …[truncated]"
        return safe_cell(value)
    return value


def _autosize(worksheet: Any, headers: list[str], rows: list[list[Any]]) -> None:
    for index, header in enumerate(headers, start=1):
        width = len(str(header))
        for row in rows:
            value = row[index - 1]
            if value is not None:
                width = max(width, min(len(str(value)), 60))
        worksheet.column_dimensions[
            worksheet.cell(row=1, column=index).column_letter
        ].width = max(10, min(width + 2, 60))


def export_xlsx(settings: Settings, output_path: Path | None = None) -> Path:
    """Write data/export_items.xlsx and return the path."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = output_path or settings.xlsx_path
    path.parent.mkdir(parents=True, exist_ok=True)

    items = items_with_summaries(settings.database_path)

    workbook = Workbook()
    items_sheet = workbook.active
    items_sheet.title = "Items"
    summaries_sheet = workbook.create_sheet("Summaries")
    bold = Font(bold=True)

    # --- Items sheet -------------------------------------------------------
    items_sheet.append(ITEM_SHEET_COLUMNS)
    for cell in items_sheet[1]:
        cell.font = bold
    item_rows: list[list[Any]] = []
    summary_rows: list[list[Any]] = []
    for item in items:
        row = [_cell_value(item.get(column)) for column in ITEM_SHEET_COLUMNS]
        item_rows.append(row)
        items_sheet.append(row)
        summary = item.get("summary") or {}
        if summary:
            summary_rows.append(
                [
                    _cell_value(item.get("id")),
                    *[
                        _cell_value(summary.get(column))
                        for column in SUMMARY_SHEET_COLUMNS[1:]
                    ],
                ]
            )
    items_sheet.freeze_panes = "A2"
    _autosize(items_sheet, ITEM_SHEET_COLUMNS, item_rows)

    # --- Summaries sheet ----------------------------------------------------
    summaries_sheet.append(SUMMARY_SHEET_COLUMNS)
    for cell in summaries_sheet[1]:
        cell.font = bold
    for row in summary_rows:
        summaries_sheet.append(row)
    summaries_sheet.freeze_panes = "A2"
    _autosize(summaries_sheet, SUMMARY_SHEET_COLUMNS, summary_rows)

    workbook.save(path)
    logger.info(
        "XLSX export written: %s (%d items, %d summaries)",
        path,
        len(item_rows),
        len(summary_rows),
    )
    return path
